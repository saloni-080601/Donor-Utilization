from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .money import format_inr, q2


NAVY = "1F5C45"
INK = "1C2A22"


def grant_master_excel(grants: list[dict], period: dict | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Grant Master"
    headers = [
        "Grant Code",
        "Donor Name",
        "Primary Location",
        "Approved Amount",
        "End Date",
        "Campus Pool",
        "Central Pool",
        "Campus Support %",
        "Central Support %",
        "Reporting cycle",
        "Next report due",
    ]
    header_fill = PatternFill("solid", fgColor=NAVY)
    header_font = Font(color="FFFFFF", bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column, header in enumerate(headers, start=1):
        cell = ws.cell(1, column, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    first_row = 2
    for offset, grant in enumerate(grants):
        row_number = first_row + offset
        values = [
            grant["code"],
            grant["donor_name"],
            grant["primary_location"],
            grant["approved_amount"],
            grant.get("end_date"),
            "Y" if grant.get("campus_pool") else "N",
            "Y" if grant.get("central_pool") else "N",
            f'=IF(F{row_number}="Y",D{row_number}/SUMIF($F${first_row}:$F${first_row + len(grants) - 1},"Y",$D${first_row}:$D${first_row + len(grants) - 1}),0)',
            f'=IF(G{row_number}="Y",D{row_number}/SUMIF($G${first_row}:$G${first_row + len(grants) - 1},"Y",$D${first_row}:$D${first_row + len(grants) - 1}),0)',
            grant.get("reporting_cycle"),
            grant.get("next_report_due"),
        ]
        for column, value in enumerate(values, start=1):
            cell = ws.cell(row_number, column, value)
            cell.alignment = Alignment(vertical="center")
            if column == 4:
                cell.number_format = '₹#,##0.00'
            elif column in {8, 9}:
                cell.number_format = "0.00%"
            elif column in {5, 11}:
                cell.number_format = "dd-mmm-yyyy"

    grant_row_count = len(grants)
    check_row = first_row + grant_row_count
    ws.cell(check_row, 1, "Check: shares add to 100%")
    ws.cell(check_row, 8, f"=ROUND(SUM(H{first_row}:H{check_row - 1}),6)" if grant_row_count else 0)
    ws.cell(check_row, 9, f"=ROUND(SUM(I{first_row}:I{check_row - 1}),6)" if grant_row_count else 0)
    for column in range(1, 12):
        ws.cell(check_row, column).fill = PatternFill("solid", fgColor="E7DFD2")
        ws.cell(check_row, column).alignment = Alignment(vertical="center")
    ws.cell(check_row, 1).font = Font(bold=True, color=INK)
    ws.cell(check_row, 8).number_format = "0.00%"
    ws.cell(check_row, 9).number_format = "0.00%"
    note_row = check_row + 2
    ws.cell(
        note_row,
        1,
        "Allocation % = approved grant / total approved grants of the grants in that pool (assumed basis, to confirm with Finance).",
    )
    ws.cell(note_row, 1).font = Font(italic=True)
    ws.cell(note_row, 1).alignment = Alignment(wrap_text=True)
    widths = [18, 24, 20, 18, 15, 14, 14, 18, 18, 18, 18]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + column)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:K{first_row + grant_row_count - 1}" if grant_row_count else "A1:K1"

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def excel_uc(grant: dict, statement: list[dict], bva: list[dict], period: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Donor UC"
    header_fill = PatternFill("solid", fgColor=NAVY)
    header_font = Font(color="FFFFFF", bold=True)
    thin = Border(
        left=Side(style="thin", color="D9CFC0"),
        right=Side(style="thin", color="D9CFC0"),
        top=Side(style="thin", color="D9CFC0"),
        bottom=Side(style="thin", color="D9CFC0"),
    )

    ws["A1"] = f"Utilization Certificate — {grant['uc_format']}"
    ws["A1"].font = Font(bold=True, size=16, color=INK)
    ws.merge_cells("A1:G1")
    meta = [
        ("Grant", f"{grant['code']} — {grant['name']}"),
        ("Donor", grant["donor_name"]),
        ("Approved", format_inr(grant["approved_amount"])),
        ("Period", f"{period['start']} to {period['end']}"),
        ("Campus pool", grant["campus_pool"]),
        ("Central pool", grant["central_pool"]),
    ]
    for i, (label, value) in enumerate(meta, start=3):
        ws[f"A{i}"] = label
        ws[f"B{i}"] = value

    start = 10
    headers = ["Ledger", "Budget head", "Direct", "Campus share", "Central share", "Total", "Overspent"]
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(start, col, title)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin
    for offset, row in enumerate(statement, start=1):
        values = [
            row["ledger"],
            row["budget_head"] or "UNMAPPED",
            q2(row["direct"]),
            q2(row["campus_alloc"]),
            q2(row["central_alloc"]),
            q2(row["total"]),
            "Yes" if row.get("overspent") else "",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(start + offset, col, value)
            cell.border = thin
            if isinstance(value, Decimal):
                cell.number_format = "#,##,##0.00"
                cell.alignment = Alignment(horizontal="right")

    bva_start = start + len(statement) + 3
    ws.cell(bva_start, 1, "Budget vs actual (grouped by head)")
    ws.cell(bva_start, 1).font = Font(bold=True)
    bva_headers = ["Budget head", "Budget", "Actual", "Variance", "%", "Overspent"]
    for col, title in enumerate(bva_headers, start=1):
        cell = ws.cell(bva_start + 1, col, title)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin
    for offset, row in enumerate(bva, start=1):
        values = [
            row["head"],
            q2(row["budget"]),
            q2(row["actual"]),
            q2(row["variance"]),
            None if row["pct"] is None else float(row["pct"]),
            "Yes" if row["overspent"] else "",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(bva_start + 1 + offset, col, value)
            cell.border = thin
            if isinstance(value, Decimal):
                cell.number_format = "#,##,##0.00"

    for col in range(1, 8):
        ws.column_dimensions[chr(64 + col)].width = 22

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pdf_uc(grant: dict, statement: list[dict], bva: list[dict], period: dict) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title=f"UC {grant['code']}")
    styles = getSampleStyleSheet()
    title = ParagraphStyle("TitleUC", parent=styles["Title"], textColor=colors.HexColor("#1F5C45"), fontSize=16)
    body = ParagraphStyle("BodyUC", parent=styles["Normal"], leading=14)
    story = [
        Paragraph(f"Utilization Certificate — {grant['uc_format']}", title),
        Paragraph(f"<b>{grant['code']}</b> · {grant['name']}<br/>Donor: {grant['donor_name']}<br/>Period: {period['start']} to {period['end']}<br/>Approved: {format_inr(grant['approved_amount'])}", body),
        Spacer(1, 12),
        Paragraph("Ledger statement", styles["Heading3"]),
    ]

    table_data = [["Ledger", "Head", "Direct", "Campus", "Central", "Total"]]
    for row in statement:
        table_data.append(
            [
                row["ledger"],
                row["budget_head"] or "UNMAPPED",
                format_inr(row["direct"]),
                format_inr(row["campus_alloc"]),
                format_inr(row["central_alloc"]),
                format_inr(row["total"]),
            ]
        )
    table = Table(table_data, repeatRows=1, colWidths=[110, 90, 70, 70, 70, 80])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F5C45")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D9CFC0")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4EFE4")]),
            ]
        )
    )
    story.extend([table, Spacer(1, 14), Paragraph("Budget vs actual", styles["Heading3"])])

    bva_data = [["Head", "Budget", "Actual", "Variance", "Status"]]
    for row in bva:
        bva_data.append(
            [
                row["head"],
                format_inr(row["budget"]),
                format_inr(row["actual"]),
                format_inr(row["variance"]),
                "Overspent" if row["overspent"] else "OK",
            ]
        )
    bva_table = Table(bva_data, repeatRows=1)
    bva_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1C2A22")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("ALIGN", (1, 1), (-2, -1), "RIGHT"),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D9CFC0")),
            ]
        )
    )
    story.extend(
        [
            bva_table,
            Spacer(1, 24),
            Paragraph("This certificate is generated from a published (or current) run. Figures are INR, rounded to the rupee before comparison. Tally is never written back to.", body),
        ]
    )
    doc.build(story)
    return buf.getvalue()
