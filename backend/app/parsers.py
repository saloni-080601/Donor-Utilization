from __future__ import annotations

import csv
import io
import calendar
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook


def _norm(value: str) -> str:
    text = str(value or "").strip().replace("_", " ").strip("*` ")
    return " ".join(text.lower().split())


def _to_decimal(value) -> Decimal:
    if value is None or value == "":
        return Decimal("0")
    text = str(value).replace(",", "").replace("₹", "").strip()
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"Not a rupee amount: {value}") from exc


def _grid_from_upload(filename: str, data: bytes) -> list[tuple]:
    name = filename.lower()
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig")
        return list(csv.reader(io.StringIO(text)))
    if name.endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(data), data_only=True)
        sheet = workbook.active
        return list(sheet.iter_rows(values_only=True))
    raise ValueError("Use a .csv or .xlsx file.")


def _header_index(rows: list[tuple]) -> int:
    for idx, row in enumerate(rows):
        values = {_norm(value) for value in row if value is not None and str(value).strip()}
        grant_header = any(value.startswith("grant") for value in values)
        approved_header = bool(values.intersection({"approved grant", "approved amount"}))
        if (
            "particulars" in values
            or {"grant", "ledger"}.issubset(values)
            or {"ledger", "pool"}.issubset(values)
            or {"code", "name"}.issubset(values)
            or (grant_header and approved_header)
        ):
            return idx
    return 0


def _rows_from_upload(filename: str, data: bytes) -> list[dict]:
    rows = _grid_from_upload(filename, data)
    if not rows:
        return []

    # Tally exports often have a title and blank lines above the actual header.
    header_idx = _header_index(rows)
    headers = [str(h or "").strip() for h in rows[header_idx]]
    out = []
    for row in rows[header_idx + 1 :]:
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        out.append(
            {
                headers[i]: row[i] if i < len(row) else None
                for i in range(len(headers))
                if headers[i]
            }
        )
    return out


def _to_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%d-%b-%Y", "%d-%B-%Y", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    for fmt in ("%b-%Y", "%B-%Y"):
        try:
            parsed = datetime.strptime(text, fmt)
            return date(parsed.year, parsed.month, calendar.monthrange(parsed.year, parsed.month)[1])
        except ValueError:
            pass
    raise ValueError(f"Not a supported date: {value}. Use YYYY-MM-DD.")


def _period_start_from_upload(filename: str, data: bytes) -> date | None:
    grid = _grid_from_upload(filename, data)
    for row in grid:
        if len(row) > 1 and _norm(row[0]) in {"period start", "start date"} and row[1] not in (None, ""):
            return _to_date(row[1])
    return None


def _pick(row: dict, *aliases: str):
    lookup = {_norm(k): v for k, v in row.items() if k}
    for alias in aliases:
        if alias in lookup and lookup[alias] not in (None, ""):
            return lookup[alias]
    raise KeyError("Missing column. Tried: " + ", ".join(aliases))


def _pick_optional(row: dict, *aliases: str):
    lookup = {_norm(k): v for k, v in row.items() if k}
    for alias in aliases:
        if alias in lookup:
            return lookup[alias]
    return None


def parse_grants(filename: str, data: bytes) -> list[dict]:
    raw_rows = _rows_from_upload(filename, data)
    if not raw_rows:
        raise ValueError("Grant Master file has no grant rows.")

    global_start = _period_start_from_upload(filename, data)
    grants = []
    for row_number, raw in enumerate(raw_rows, start=2):
        code = _pick_optional(raw, "code", "grant code", "grant", "grant as in tally export", "grant (as in tally export)", "grant name")
        if code is None or not str(code).strip():
            continue
        code = str(code).strip()
        if _norm(code).startswith("other grants in central pool"):
            code = "OTHER-CENTRAL"
        name = str(
            _pick_optional(raw, "name", "grant display name")
            or _pick_optional(raw, "grant", "grant as in tally export", "grant (as in tally export)")
            or code
        ).strip()
        donor = str(_pick(raw, "donor name", "donor", "donor_name")).strip()
        amount = _to_decimal(_pick(raw, "approved amount", "approved grant", "approved_amount", "amount"))
        start_value = _pick_optional(raw, "start date", "start_date") or global_start
        if start_value is None:
            raise ValueError(f"Grant row {row_number} needs a start_date or a Period start above the table.")
        end_value = _pick_optional(raw, "end date", "end_date", "grant end date")
        location_value = _pick_optional(raw, "primary location", "location", "primary_location")
        location = str(location_value or ("Other" if code == "OTHER-CENTRAL" else "")).strip()
        if not location:
            raise ValueError(f"Grant row {row_number} needs a primary location.")

        campus_pool = _pick_optional(raw, "campus pool", "campus_pool")
        if campus_pool in (None, ""):
            campus_flag = str(_pick_optional(raw, "in campus pool", "in campus pool (y/n)") or "").strip().lower()
            campus_pool = "CAMPUS-SUPPORT" if campus_flag in {"y", "yes", "true", "1"} else ""
        central_pool = _pick_optional(raw, "central pool", "central_pool")
        if central_pool in (None, ""):
            central_flag = str(_pick_optional(raw, "in central pool", "in central pool (y/n)") or "").strip().lower()
            central_pool = "CENTRAL-SUPPORT" if central_flag in {"y", "yes", "true", "1"} else ""

        grants.append(
            {
                "code": code,
                "name": name,
                "donor_name": donor,
                "uc_format": str(_pick_optional(raw, "uc format", "uc_format") or "Standard UC").strip(),
                "approved_amount": amount,
                "start_date": _to_date(start_value),
                "end_date": _to_date(end_value) if end_value not in (None, "") else None,
                "campus_pool": str(campus_pool or "").strip(),
                "central_pool": str(central_pool or "").strip(),
                "primary_location": location,
            }
        )
    if not grants:
        raise ValueError("Grant Master file has no grant rows.")
    return grants


def parse_utilization(filename: str, data: bytes) -> list[dict]:
    raw_rows = _rows_from_upload(filename, data)
    if not raw_rows:
        raise ValueError("Utilization file has no spend rows.")

    headers = {_norm(key) for key in raw_rows[0]}
    rows = []
    if headers.intersection({"grant", "grant code", "grant name", "grantcode"}):
        for raw in raw_rows:
            grant = str(_pick(raw, "grant", "grant code", "grant name", "grantcode")).strip()
            ledger = str(_pick(raw, "ledger", "ledger name", "ledgername")).strip()
            amount = _to_decimal(_pick(raw, "amount", "spend", "debit", "utilization"))
            if amount != 0:
                rows.append({"grant_code": grant, "ledger": ledger, "amount": amount})
    elif "particulars" in headers:
        excluded = {"particulars", "total", "grand total", "general", "unfunded", "campus support", "central support", "budget head sr no"}
        for raw in raw_rows:
            ledger = str(_pick(raw, "particulars")).strip()
            if not ledger or _norm(ledger).startswith("grand total") or _norm(ledger).startswith("allocation %"):
                continue
            for grant_column, value in raw.items():
                grant = str(grant_column).strip()
                if _norm(grant) in excluded or _norm(grant).startswith("budget head"):
                    continue
                amount = _to_decimal(value)
                if amount != 0:
                    rows.append({"grant_code": grant, "ledger": ledger, "amount": amount})
    else:
        _pick(raw_rows[0], "grant", "grant code", "grant name", "grantcode")
    if not rows:
        raise ValueError("Utilization file has no spend rows.")
    return rows


def parse_central(filename: str, data: bytes) -> list[dict]:
    raw_rows = _rows_from_upload(filename, data)
    headers = {_norm(key) for key in raw_rows[0]} if raw_rows else set()
    rows = []
    if headers.intersection({"pool", "cost pool", "pool name"}):
        for raw in raw_rows:
            ledger = str(_pick(raw, "ledger", "ledger name", "ledgername")).strip()
            pool = str(_pick(raw, "pool", "cost pool", "pool name")).strip()
            kind_raw = str(_pick(raw, "pool type", "pool kind", "type", "kind")).strip().lower()
            if kind_raw in {"campus", "campus pool"}:
                kind = "campus"
            elif kind_raw in {"central", "central pool"}:
                kind = "central"
            else:
                raise ValueError(f"Pool type must be campus or central, got {kind_raw}")
            amount = _to_decimal(_pick(raw, "amount", "spend", "cost"))
            if amount != 0:
                rows.append({"ledger": ledger, "pool": pool, "pool_kind": kind, "amount": amount})
    elif "particulars" in headers:
        pool_columns = {
            "campus project support": ("__AUTO_CAMPUS_POOL__", "campus"),
            "campus support": ("__AUTO_CAMPUS_POOL__", "campus"),
            "central support": ("__AUTO_CENTRAL_POOL__", "central"),
        }
        for raw in raw_rows:
            ledger = str(_pick(raw, "particulars")).strip()
            if not ledger or _norm(ledger).startswith("grand total"):
                continue
            for column, (pool, kind) in pool_columns.items():
                value = next((v for k, v in raw.items() if _norm(k) == column), None)
                amount = _to_decimal(value)
                if amount != 0:
                    rows.append({"ledger": ledger, "pool": pool, "pool_kind": kind, "amount": amount})
    else:
        _pick(raw_rows[0], "pool", "cost pool", "pool name")
    if not rows:
        raise ValueError("Central costs file has no rows.")
    return rows
